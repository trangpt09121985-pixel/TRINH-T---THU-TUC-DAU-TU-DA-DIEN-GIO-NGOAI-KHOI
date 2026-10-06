# -*- coding: utf-8 -*-
"""Công cụ cập nhật dữ liệu cho dashboard ĐGNK PVEP.

Cách dùng (chạy từ thư mục gốc của dự án):
  python tools/build_data.py                    # đọc data/*.csv  -> sinh data/data.js
  python tools/build_data.py --from-xlsx FILE   # đọc file Excel  -> ghi lại data/*.csv và data/data.js
  python tools/build_data.py --to-xlsx FILE     # xuất data/*.csv -> file Excel mẫu để cập nhật
Tùy chọn thêm:
  --data THU_MUC   thư mục dữ liệu (mặc định: data/ của bộ mã)
  --bo noi_bo      đánh dấu bộ dữ liệu nội bộ (mặc định: cong_khai)

Mỗi bảng tương ứng một tệp CSV (UTF-8) hoặc một trang tính cùng tên trong Excel.
Công cụ kiểm tra cột bắt buộc, mã trùng và mã văn bản không tồn tại trước khi ghi.
"""
import csv, json, os, sys, datetime

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, 'data')
TABLES = {
    'thu_tuc':          ('thu_tuc.csv', 'ma'),
    'cong_quyet_dinh':  ('cong_quyet_dinh.csv', 'ma'),
    'van_ban':          ('van_ban.csv', 'ma_vb'),
    'ma_tran_phap_ly':  ('ma_tran_phap_ly.csv', 'ma'),
    'rui_ro':           ('rui_ro.csv', 'ma'),
    'ke_hoach':         ('ke_hoach.csv', 'ma'),
    'giai_doan':        ('giai_doan.csv', 'ma'),
}
LABELS = {'ĐÃ XÁC MINH', 'CẦN XÁC MINH', 'CHƯA CÓ HƯỚNG DẪN CHI TIẾT', 'CẦN Ý KIẾN CƠ QUAN CÓ THẨM QUYỀN'}
ALLOWED = {
    ('thu_tuc', 'nhan_xac_minh'): LABELS | {'NGHIỆP VỤ NỘI BỘ'},
    ('thu_tuc', 'luong'): {'A', 'B', 'Chung'},
    ('thu_tuc', 'pham_vi'): {'Ngoài khơi', 'Trên bờ', 'Chung'},
    ('van_ban', 'trang_thai_xac_minh'): LABELS | {'THAM KHẢO KỸ THUẬT'},
    ('ma_tran_phap_ly', 'trang_thai_xac_minh'): LABELS,
    ('thu_tuc', 'muc_rui_ro'): {'Cao', 'Trung bình', 'Thấp'},
    ('thu_tuc', 'thoi_diem'): {'Hiện tại', 'Chưa đến thời điểm'},
    ('thu_tuc', 'trang_thai_thuc_hien'): {'Chưa bắt đầu', 'Đang thực hiện', 'Hoàn thành', 'Tạm dừng'},
    ('thu_tuc', 'uu_tien'): {'Cao', 'Trung bình', 'Thấp'},
    ('thu_tuc', 'can_xin_y_kien'): {'Có', 'Không'},
    ('ma_tran_phap_ly', 'loai'): {'Quy định đã rõ', 'Khoảng trống'},
    ('ke_hoach', 'tinh_trang'): {'Chưa bắt đầu', 'Đang thực hiện', 'Hoàn thành', 'Tạm dừng'},
}

def read_csv(path):
    with open(path, encoding='utf-8-sig', newline='') as f:
        return [{k.strip(): (v or '').strip() for k, v in r.items()} for r in csv.DictReader(f)]

def write_csv(path, rows):
    if not rows:
        return
    with open(path, 'w', encoding='utf-8-sig', newline='') as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        w.writeheader(); w.writerows(rows)

def validate(db):
    errs = []
    vb_ids = {r['ma_vb'] for r in db['van_ban']}
    for name, (_, key) in TABLES.items():
        seen = set()
        for i, r in enumerate(db[name], start=2):
            k = r.get(key, '')
            if not k:
                errs.append(f'{name} dòng {i}: thiếu mã ({key})')
            elif k in seen:
                errs.append(f'{name} dòng {i}: mã trùng {k}')
            seen.add(k)
            for (t, col), allowed in ALLOWED.items():
                if t == name and r.get(col, '') and r[col] not in allowed:
                    errs.append(f'{name} dòng {i} ({k}): giá trị “{r[col]}” ở cột {col} không hợp lệ')
    for r in db['thu_tuc']:
        for m in filter(None, (x.strip() for x in r.get('ma_van_ban', '').split(';'))):
            if m not in vb_ids:
                errs.append(f'thu_tuc {r["ma"]}: mã văn bản {m} không có trong van_ban')
    tt_ids = {r['ma'] for r in db['thu_tuc']}
    for r in db['thu_tuc']:
        for col in ('dieu_kien_tien_quyet', 'song_song_voi'):
            for m in filter(None, (x.strip() for x in r.get(col, '').split(';'))):
                if m not in tt_ids:
                    errs.append(f'thu_tuc {r["ma"]}: mã {m} ở cột {col} không tồn tại')
                if m == r['ma']:
                    errs.append(f'thu_tuc {r["ma"]}: cột {col} tham chiếu chính nó')
    ph_ids = {r['ma'] for r in db['giai_doan']}
    for r in db['thu_tuc']:
        if r.get('giai_doan_10', '') not in ph_ids:
            errs.append(f'thu_tuc {r["ma"]}: giai_doan_10 “{r.get("giai_doan_10", "")}” không có trong bảng giai_doan')
    gate_ids = {r['ma'] for r in db['cong_quyet_dinh']}
    for r in db['giai_doan']:
        if r.get('cong_quyet_dinh') and r['cong_quyet_dinh'] not in gate_ids:
            errs.append(f'giai_doan {r["ma"]}: cổng {r["cong_quyet_dinh"]} không tồn tại')
    for r in db['cong_quyet_dinh']:
        for m in filter(None, (x.strip() for x in r.get('thu_tuc_dau_vao', '').split(';'))):
            if m not in tt_ids:
                errs.append(f'cong_quyet_dinh {r["ma"]}: mã đầu vào {m} không tồn tại')
    for r in db['van_ban']:
        if r.get('cap_van_ban', '') not in ('', '1', '2', '3', '4', '5'):
            errs.append(f'van_ban {r["ma_vb"]}: cap_van_ban phải là 1–5')
        for col in ('quy_dinh_chi_tiet_cho', 'sua_doi_cho', 'ap_dung_uu_tien_cho'):
            for m in filter(None, (x.strip() for x in r.get(col, '').split(';'))):
                if m not in vb_ids:
                    errs.append(f'van_ban {r["ma_vb"]}: mã {m} ở cột {col} không có trong van_ban')
    for r in db['ma_tran_phap_ly']:
        if r['ma_vb'] not in vb_ids:
            errs.append(f'ma_tran_phap_ly {r["ma"]}: mã văn bản {r["ma_vb"]} không có trong van_ban')
    return errs

def load_csv_db():
    return {n: read_csv(os.path.join(DATA, f)) for n, (f, _) in TABLES.items()}

def load_xlsx_db(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    db = {}
    for n in TABLES:
        if n not in wb.sheetnames:
            raise SystemExit(f'Thiếu trang tính “{n}” trong {path}')
        ws = wb[n]
        rows = list(ws.iter_rows(values_only=True))
        head = [str(h).strip() for h in rows[0] if h is not None]
        out = []
        for row in rows[1:]:
            vals = ['' if v is None else (v.strftime('%Y-%m-%d') if isinstance(v, (datetime.date, datetime.datetime)) else str(v).strip()) for v in row[:len(head)]]
            if any(vals):
                out.append(dict(zip(head, vals)))
        db[n] = out
    return db

BO = 'cong_khai'
def write_js(db):
    meta = {'ngay_tao': datetime.datetime.now().strftime('%Y-%m-%d %H:%M'), 'phien_ban': datetime.date.today().isoformat(), 'bo_du_lieu': BO}
    payload = json.dumps({'meta': meta, **db}, ensure_ascii=False, indent=1)
    with open(os.path.join(DATA, 'data.js'), 'w', encoding='utf-8') as f:
        f.write('/* Tệp sinh tự động bởi tools/build_data.py – không sửa trực tiếp. */\nwindow.APP_DATA = ' + payload + ';\n')

def to_xlsx(db, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook(); ws0 = wb.active; ws0.title = 'Huong_dan'
    guide = [
        'HƯỚNG DẪN CẬP NHẬT DỮ LIỆU – BẢN ĐỒ THỦ TỤC ĐGNK PVEP',
        '1. Mỗi trang tính là một bảng dữ liệu; không đổi tên trang tính, không đổi tên cột ở dòng 1.',
        '2. Cột mã (ma, ma_vb) là duy nhất. Thêm dòng mới ở cuối bảng.',
        '3. Các cột có danh sách thả xuống chỉ nhận giá trị trong danh sách.',
        '4. Cột ma_van_ban ở bảng thu_tuc: ghi các mã VBxx cách nhau bằng dấu chấm phẩy; mã phải có trong bảng van_ban. Cột dieu_kien_tien_quyet, song_song_voi: ghi mã TT-x.y cách nhau bằng dấu chấm phẩy.',
        '4a. Nhãn xác minh chỉ dùng: ĐÃ XÁC MINH, CẦN XÁC MINH, CHƯA CÓ HƯỚNG DẪN CHI TIẾT, CẦN Ý KIẾN CƠ QUAN CÓ THẨM QUYỀN (và NGHIỆP VỤ NỘI BỘ cho việc thuần nội bộ).',
        '5. Không điền kết luận pháp lý, thời hạn, thẩm quyền khi chưa có văn bản gốc; ghi “CẦN XÁC MINH”.',
        '6. Ngày ghi theo định dạng YYYY-MM-DD.',
        '7. Sau khi cập nhật: python tools/build_data.py --from-xlsx <tên tệp> để sinh lại data/data.js. Bộ nội bộ: thêm --data <thư mục> --bo noi_bo.',
        'Dòng tô vàng ở mỗi bảng là dòng dữ liệu mẫu đầu tiên, thể hiện định dạng mong đợi.',
    ]
    for i, t in enumerate(guide, 1):
        c = ws0.cell(row=i, column=1, value=t); c.font = Font(name='Arial', bold=(i == 1), size=12 if i == 1 else 11)
    ws0.column_dimensions['A'].width = 120
    head_fill = PatternFill('solid', fgColor='005088'); ex_fill = PatternFill('solid', fgColor='FFF4CC')
    for n in TABLES:
        rows = db[n]; ws = wb.create_sheet(n)
        cols = list(rows[0].keys())
        for j, c in enumerate(cols, 1):
            cell = ws.cell(row=1, column=j, value=c)
            cell.font = Font(name='Arial', bold=True, color='FFFFFF'); cell.fill = head_fill
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            ws.column_dimensions[cell.column_letter].width = 16 if c in ('ma', 'ma_vb', 'nhanh') else 34
        for i, r in enumerate(rows, 2):
            for j, c in enumerate(cols, 1):
                cell = ws.cell(row=i, column=j, value=r[c])
                cell.font = Font(name='Arial', size=10); cell.alignment = Alignment(wrap_text=True, vertical='top')
                if i == 2:
                    cell.fill = ex_fill
        ws.freeze_panes = 'B2'
        for (t, col), allowed in ALLOWED.items():
            if t == n and col in cols:
                letter = ws.cell(row=1, column=cols.index(col) + 1).column_letter
                dv = DataValidation(type='list', formula1='"' + ','.join(sorted(allowed)) + '"', allow_blank=True)
                ws.add_data_validation(dv); dv.add(f'{letter}2:{letter}1000')
    wb.save(path)

def main():
    global DATA, BO
    args = sys.argv[1:]
    if '--data' in args:
        i = args.index('--data'); DATA = os.path.abspath(args[i + 1]); del args[i:i + 2]
    if '--bo' in args:
        i = args.index('--bo'); BO = args[i + 1]; del args[i:i + 2]
    if args[:1] == ['--to-xlsx']:
        to_xlsx(load_csv_db(), args[1]); print('Đã xuất', args[1]); return
    db = load_xlsx_db(args[1]) if args[:1] == ['--from-xlsx'] else load_csv_db()
    errs = validate(db)
    if errs:
        print('DỮ LIỆU CHƯA HỢP LỆ – không ghi tệp:'); [print(' -', e) for e in errs]; sys.exit(1)
    if args[:1] == ['--from-xlsx']:
        for n, (f, _) in TABLES.items():
            write_csv(os.path.join(DATA, f), db[n])
    write_js(db)
    print('Đã sinh data/data.js:', ', '.join(f'{n}={len(db[n])}' for n in TABLES))

if __name__ == '__main__':
    main()
